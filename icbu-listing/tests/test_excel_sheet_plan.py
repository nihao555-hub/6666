"""Fill sheet is a family short form; official attrs come from schema.get later."""

import os
import sys
import tempfile
import unittest
import unittest.mock
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "backend"))

os.environ["DATABASE_PATH"] = str(Path(tempfile.mkdtemp()) / "sheet-plan.db")
os.environ["UPLOAD_DIR"] = str(Path(tempfile.mkdtemp()) / "uploads")
os.environ["TOKEN_ENCRYPTION_KEY"] = "Frqkv73e3W19RfKb5WSgTueumPWcBSmpbM4wDyqu52o="
os.environ["REGISTRATION_CODES"] = "TEST-CODE"
os.environ["ALIBABA_APP_KEY"] = "test-key"
os.environ["ALIBABA_APP_SECRET"] = "test-secret"
os.environ.pop("OPENAI_API_KEY", None)

from fastapi.testclient import TestClient  # noqa: E402

from server.db import init_db  # noqa: E402
from server.main import app  # noqa: E402


def signup(email: str) -> TestClient:
    client = TestClient(app)
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "supersecret", "code": "TEST-CODE"},
    )
    assert response.status_code == 200, response.text
    return client


class SheetPlanTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        init_db()

    def test_sheet_plan_uses_family_without_category_id(self) -> None:
        client = signup("sheet-plan@example.com")
        pencils = client.get(
            "/api/v1/excel/sheet-plan",
            params={"category_name": "Office & School Supplies / Colored Pencils"},
        )
        brushes = client.get(
            "/api/v1/excel/sheet-plan",
            params={"category_name": "Tools & Hardware / Paint Brushes"},
        )
        self.assertEqual(pencils.status_code, 200, pencils.text)
        self.assertEqual(brushes.status_code, 200, brushes.text)
        pencil_plan = pencils.json()
        brush_plan = brushes.json()
        self.assertEqual(pencil_plan["sheet"]["family_id"], "stationery")
        self.assertEqual(brush_plan["sheet"]["family_id"], "tools")
        pencil_cols = [item["label"] for item in pencil_plan["preview"]["columns"]]
        brush_cols = [item["label"] for item in brush_plan["preview"]["columns"]]
        self.assertIn("色数", pencil_cols)
        self.assertIn("尺寸", brush_cols)
        self.assertNotEqual(pencil_cols, brush_cols)

    def test_sheet_plan_uses_leaf_schema_when_category_id_set(self) -> None:
        client = signup("leaf-plan@example.com")
        shop_id = client.post("/api/v1/shops/bind-env", json={"name": "schema店"}).json()["id"]
        with unittest.mock.patch(
            "server.routers.excel._attr_columns",
            return_value=[
                {
                    "id": "attr.icbuCatProp.p-2",
                    "header": "Lead Hardness",
                    "label": "铅芯硬度",
                    "group": "icbuCatProp",
                    "field_id": "p-2",
                    "required": True,
                    "options": [],
                }
            ],
        ):
            response = client.get(
                "/api/v1/excel/sheet-plan",
                params={"shop_id": shop_id, "category_id": "21110712", "category_name": "Colored Pencils"},
            )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertEqual(payload["sheet"]["family_id"], "leaf")
        self.assertEqual(payload["sheet_origin"]["kind"], "leaf_schema")
        labels = [item["label"] for item in payload["preview"]["columns"]]
        self.assertIn("铅芯硬度", labels)
        self.assertNotIn("色数", labels)

    def test_official_attrs_are_a_separate_leaf_call(self) -> None:
        client = signup("official-attrs@example.com")
        shop_id = client.post("/api/v1/shops/bind-env", json={"name": "官方属性店"}).json()["id"]
        with unittest.mock.patch("server.routers.excel._attr_columns", return_value=[
            {"id": "leadHardness", "header": "Lead Hardness", "name": "Lead Hardness"}
        ]) as mocked:
            response = client.get(
                "/api/v1/excel/official-attrs",
                params={"shop_id": shop_id, "category_id": "123"},
            )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertEqual(payload["category_id"], "123")
        self.assertEqual(payload["ai_attrs"][0]["header"], "Lead Hardness")
        self.assertTrue(any("Hardness" in item["label"] or "硬度" in item["label"] for item in payload["ai_fills"]))
        mocked.assert_called()
        _, kwargs = mocked.call_args
        self.assertTrue(kwargs.get("fetch"))

    def test_simple_template_uses_family_columns_from_category_name(self) -> None:
        client = signup("family-xlsx@example.com")
        response = client.get(
            "/api/v1/excel/template",
            params={"style": "simple", "category_name": "Office & School Supplies / Other"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        from io import BytesIO

        from openpyxl import load_workbook

        book = load_workbook(BytesIO(response.content))
        fill_sheet = next(name for name in book.sheetnames if name != "说明")
        headers = [cell.value for cell in next(book[fill_sheet].iter_rows(min_row=1, max_row=1))]
        self.assertIn("色数", headers)
        self.assertIn("硬度", headers)
        self.assertNotIn("Lead Color", headers)
        self.assertNotIn("Lead Hardness", headers)

    def test_sheet_plan_uses_full_schema_when_style_set(self) -> None:
        client = signup("full-plan@example.com")
        shop_id = client.post("/api/v1/shops/bind-env", json={"name": "完整表店"}).json()["id"]
        mocked_cols = [
            {
                "id": "schema.productTitle",
                "header": "Product Title",
                "label": "Product Title",
                "group": "productTitle",
                "field_id": "productTitle",
                "field_path": "productTitle",
                "required": True,
                "options": [],
            },
            {
                "id": "schema.icbuCatProp.p-type",
                "header": "icbuCatProp / Type",
                "label": "Type",
                "group": "icbuCatProp",
                "field_id": "p-type",
                "field_path": "icbuCatProp.p-type",
                "required": True,
                "options": [],
            },
            {
                "id": "schema.icbuCatProp.p-color",
                "header": "icbuCatProp / Color",
                "label": "Color",
                "group": "icbuCatProp",
                "field_id": "p-color",
                "field_path": "icbuCatProp.p-color",
                "required": False,
                "options": [],
            },
        ]
        with unittest.mock.patch("server.routers.excel._schema_columns", return_value=mocked_cols):
            response = client.get(
                "/api/v1/excel/sheet-plan",
                params={
                    "shop_id": shop_id,
                    "category_id": "21110712",
                    "category_name": "Colored Pencils",
                    "style": "full_schema",
                },
            )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertEqual(payload["sheet"]["family_id"], "full_schema")
        self.assertEqual(payload["sheet_origin"]["kind"], "full_schema")
        self.assertEqual(payload["sheet"]["required_count"], 2)
        self.assertEqual(payload["sheet"]["optional_count"], 1)
        labels = [item["label"] for item in payload["preview"]["columns"]]
        self.assertIn("英文标题", labels)
        self.assertIn("icbuCatProp / Type", labels)
        self.assertIn("icbuCatProp / Color", labels)


if __name__ == "__main__":
    unittest.main()

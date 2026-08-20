"""Excel styles, header detection, and official templates."""

import io
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "backend"))

from openpyxl import Workbook, load_workbook  # noqa: E402

from schema import parse_schema  # noqa: E402
from server.services.excel_import import (  # noqa: E402
    ExcelRow,
    apply_preview,
    build_template,
    category_attr_columns,
    decide_image_action,
    fill_policy,
    find_header_row,
    flatten_schema_fields,
    guess_field,
    guess_style,
    mapping_from_headers,
    match_uploads,
    parse_rows,
    preview,
    resolve_row_files,
    schema_field_columns,
    sheet_preview,
    sheet_profile,
    split_images,
)


class AliasTests(unittest.TestCase):
    def test_lingxing_and_mabang_headers_map(self) -> None:
        self.assertEqual(guess_field("货号"), "sku")
        self.assertEqual(guess_field("单价USD"), "price")
        self.assertEqual(guess_field("库存SKU", "mabang"), "sku")
        self.assertEqual(guess_field("中文名称", "mabang"), "name")
        self.assertEqual(guess_field("图片地址", "mabang"), "images")

    def test_style_guess_from_headers(self) -> None:
        self.assertEqual(guess_style(["库存SKU", "中文名称", "售价"]), "mabang")
        self.assertEqual(guess_style(["货号", "单价USD", "起订量"]), "lingxing")
        self.assertEqual(guess_style(["货号", "单价 USD", "起订量", "图片", "品牌"]), "simple")
        self.assertEqual(guess_style(["货号", "英文标题", "关键词"]), "dianxiaomi")


class ParseTests(unittest.TestCase):
    def test_skips_title_rows_to_find_the_real_header(self) -> None:
        rows = [
            ["马帮库存导出", "", ""],
            ["导出时间", "2026-08-16", ""],
            ["库存SKU", "售价", "图片地址"],
            ["BRUSH-01", "1.8", "a.jpg;b.jpg"],
        ]
        index = find_header_row(rows, "mabang")
        self.assertEqual(index, 2)
        mapping = mapping_from_headers(rows[index], "mabang")
        parsed = parse_rows(rows, mapping, index)
        self.assertEqual(len(parsed), 1)
        self.assertEqual(parsed[0].sku, "BRUSH-01")
        self.assertEqual(parsed[0].price, "1.8")
        self.assertEqual(parsed[0].images, ["a.jpg", "b.jpg"])

    def test_excel_row_marks_seed_fields_as_excel(self) -> None:
        row = ExcelRow(title="Paint Brush Set", price="1.80", moq="500", origin="Vietnam")
        self.assertEqual(row.seed_values()["productTitle"], "Paint Brush Set")
        self.assertEqual(row.provided_sources()["productTitle"], "excel")
        self.assertEqual(row.provided_sources()["ladderPrice"], "excel")
        self.assertEqual(row.provided_sources()["origin"], "excel")

    def test_image_refs_split_urls_from_filenames(self) -> None:
        urls, names = split_images(["https://img.example.com/a.jpg", "SKU-1_1.jpg", "not a url"])
        self.assertEqual(urls, ["https://img.example.com/a.jpg"])
        self.assertEqual(names, ["SKU-1_1.jpg", "not a url"])

    def test_uploads_match_sku_prefix(self) -> None:
        uploads = {
            "sku-1001_1.jpg": b"one",
            "sku-1001-2.jpg": b"two",
            "other.jpg": b"nope",
        }
        matched = match_uploads("SKU-1001", [], uploads)
        self.assertEqual({name for name, _ in matched}, {"sku-1001_1.jpg", "sku-1001-2.jpg"})

    def test_one_sku_named_file_is_enough(self) -> None:
        matched = match_uploads("SKU-1001", [], {"sku-1001.jpg": b"only"})
        self.assertEqual(matched, [("sku-1001.jpg", b"only")])
        files = resolve_row_files(ExcelRow(sku="SKU-1001"), {"sku-1001.jpg": b"only"})
        self.assertEqual(len(files), 1)


class TemplateTests(unittest.TestCase):
    def test_sheet_preview_is_the_short_form_not_official_forty(self) -> None:
        preview_data = sheet_preview("simple")
        labels = [item["label"] for item in preview_data["columns"]]
        self.assertEqual(labels, ["货号", "单价 USD", "起订量", "图片", "品牌", "品名（中文）", "备注"])
        self.assertFalse(preview_data["from_official_form"])
        self.assertNotIn("英文标题", labels)
        self.assertNotIn("叶子类目 ID", labels)
        self.assertIn("不是阿里后台", preview_data["note"])
        images_col = next(item for item in preview_data["columns"] if item["id"] == "images")
        self.assertFalse(images_col["required"])
        self.assertTrue(next(item for item in preview_data["columns"] if item["id"] == "price")["required"])

    def test_generated_lingxing_template_round_trips(self) -> None:
        payload = build_template("lingxing")
        result = preview(payload, "lingxing")
        self.assertEqual(result["mapping"]["货号"], "sku")
        self.assertEqual(result["mapping"]["单价 USD"], "price")
        # An untouched template holds nothing but the worked example.
        self.assertEqual(result["row_count"], 0)
        self.assertEqual(result["sample_skipped"], 1)
        self.assertEqual(apply_preview(payload, result["mapping"], "lingxing"), [])

    def test_required_attribute_columns_differ_by_category(self) -> None:
        brushes = parse_schema(
            """<?xml version="1.0"?><itemSchema>
              <field id="icbuCatProp" type="complex"><fields>
                <field id="p-1" name="Place of Origin" type="singleCheck"><rules><rule name="requiredRule" value="true"/></rules></field>
                <field id="p-type" name="Type" type="singleCheck"><rules><rule name="requiredRule" value="true"/></rules>
                  <options><option displayName="Oil Brush" value="1"/></options>
                </field>
              </fields></field>
              <field id="saleProp" type="complex"><fields>
                <field id="p-color" name="Color" type="singleCheck"><rules><rule name="requiredRule" value="true"/></rules>
                  <options><option displayName="Black" value="9"/></options>
                </field>
              </fields></field>
            </itemSchema>"""
        )
        pens = parse_schema(
            """<?xml version="1.0"?><itemSchema>
              <field id="icbuCatProp" type="complex"><fields>
                <field id="p-1" name="Place of Origin" type="singleCheck"><rules><rule name="requiredRule" value="true"/></rules></field>
                <field id="p-hair" name="Hair Material" type="singleCheck"><rules><rule name="requiredRule" value="true"/></rules>
                  <options><option displayName="Wolf Hair" value="2"/></options>
                </field>
              </fields></field>
            </itemSchema>"""
        )
        brush_cols = {item["header"] for item in category_attr_columns(brushes)}
        pen_cols = {item["header"] for item in category_attr_columns(pens)}
        self.assertEqual(brush_cols, {"Type", "Color"})
        self.assertEqual(pen_cols, {"Hair Material"})
        self.assertNotIn("Place of Origin", brush_cols)
        extras = category_attr_columns(brushes)
        profile = sheet_profile("Paint Brushes", category_id="21111112", attr_columns=extras)
        payload = build_template(
            "simple",
            {"name": "Paint Brushes", "category_id": "21111112"},
            extra_columns=extras,
            category_id="21111112",
        )
        result = preview(payload, "simple", extras)
        self.assertIn("类型", result["headers"])
        self.assertIn("颜色", result["headers"])
        self.assertNotIn("英文标题", result["headers"])
        book = load_workbook(io.BytesIO(payload))
        help_text = " ".join(
            str(cell or "")
            for row in book["说明"].iter_rows(values_only=True)
            for cell in row
        )
        self.assertIn("类型", help_text)
        self.assertIn("颜色", help_text)
        self.assertIn("红线", help_text)
        policy = fill_policy(extras, profile)
        user_labels = {item["label"] for item in policy["user_fills"]}
        self.assertIn("类型", user_labels)
        self.assertIn("颜色", user_labels)
        self.assertTrue(any(item["id"] == "price" for item in policy["redline"]))

    def test_uploaded_official_attr_columns_still_parse(self) -> None:
        extras = category_attr_columns(
            parse_schema(
                """<?xml version="1.0"?><itemSchema>
                  <field id="icbuCatProp" type="complex"><fields>
                    <field id="p-type" name="Type" type="singleCheck"><rules><rule name="requiredRule" value="true"/></rules>
                      <options><option displayName="Oil Brush" value="1"/></options>
                    </field>
                  </fields></field>
                </itemSchema>"""
            )
        )
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.append(["货号", "单价 USD", "起订量", "图片", "Type"])
        sheet.append(["SKU-1001", "1.80", "500", "a.jpg", "Oil Brush"])
        buffer = io.BytesIO()
        book.save(buffer)
        rows = apply_preview(buffer.getvalue(), {"货号": "sku", "单价 USD": "price", "Type": "attr.icbuCatProp.p-type"}, "simple", extras)
        self.assertEqual(rows[0].seed_values()["icbuCatProp"]["p-type"], "1")

    def test_simple_template_is_only_what_the_seller_must_fill(self) -> None:
        payload = build_template("simple")
        result = preview(payload, "simple")
        self.assertEqual(set(result["mapping"].values()), {"sku", "name", "price", "moq", "images", "note", "brand"})
        self.assertNotIn("英文标题", result["headers"])
        self.assertNotIn("叶子类目 ID", result["headers"])
        book = load_workbook(io.BytesIO(payload))
        help_text = " ".join(
            str(cell or "")
            for row in book["说明"].iter_rows(values_only=True)
            for cell in row
        )
        self.assertIn("补转化位", help_text)
        images_fill = next(item for item in fill_policy()["user_fills"] if item["id"] == "images")
        self.assertFalse(images_fill["required"])

    def test_trade_and_logistics_are_shop_defaults_not_ai(self) -> None:
        policy = fill_policy()
        self.assertNotIn("交易和物流", {item["label"] for item in policy["ai_fills"]})
        self.assertIn("运费模板", {item["label"] for item in policy["shop_fills"]})
        self.assertIn("计量单位", {item["label"] for item in policy["shop_fills"]})
        self.assertIn("零出错", policy["guarantee"])
        note = next(item for item in policy["user_fills"] if item["id"] == "note")
        self.assertIn("规格列", note["hint"])
        book = load_workbook(io.BytesIO(build_template("simple")))
        help_text = " ".join(str(cell or "") for row in book["说明"].iter_rows(values_only=True) for cell in row)
        self.assertIn("店里套", help_text)
        self.assertIn("人核对", help_text)

    def test_one_sheet_carries_many_products(self) -> None:
        payload = build_template("simple")
        book = load_workbook(io.BytesIO(payload))
        sheet = book["填写"]
        for index in range(1, 4):
            sheet.append([f"A-{index}", "2.30", "300", f"A-{index}_1.jpg", "", f"毛笔{index}", ""])
        buffer = io.BytesIO()
        book.save(buffer)

        result = preview(buffer.getvalue(), "simple")
        self.assertEqual(result["row_count"], 3)
        self.assertEqual(result["ready_count"], 3)
        rows = apply_preview(buffer.getvalue(), result["mapping"], "simple")
        self.assertEqual([row.sku for row in rows], ["A-1", "A-2", "A-3"])

    def test_worked_example_never_becomes_a_product(self) -> None:
        """Sellers append below the sample as often as they overwrite it."""
        payload = build_template("simple")
        book = load_workbook(io.BytesIO(payload))
        sheet = book["填写"]
        sheet.append(["A-01", "2.30", "300", "A-01_1.jpg", "", "毛笔", ""])
        buffer = io.BytesIO()
        book.save(buffer)

        result = preview(buffer.getvalue(), "simple")
        self.assertEqual(result["row_count"], 1)
        self.assertEqual(result["sample_skipped"], 1)
        rows = apply_preview(buffer.getvalue(), result["mapping"], "simple")
        self.assertEqual([row.sku for row in rows], ["A-01"])
        self.assertTrue(any("示例" in warning for warning in result["warnings"]))

    def test_rows_are_checked_before_any_ai_time_is_spent(self) -> None:
        book = Workbook()
        sheet = book.active
        sheet.append(["货号", "单价 USD", "起订量", "图片"])
        sheet.append(["A-01", "2.30", "300", "A-01_1.jpg"])
        sheet.append(["A-02", "", "300", "A-02_1.jpg"])
        sheet.append(["A-03", "1.00", "", ""])
        sheet.append(["A-01", "1.00", "100", "dup.jpg"])
        buffer = io.BytesIO()
        book.save(buffer)

        result = preview(buffer.getvalue(), "simple")
        self.assertEqual(result["row_count"], 4)
        self.assertEqual(result["blocked_count"], 2)
        self.assertEqual(result["ready_count"], 2)
        messages = {(item["line"], item["level"]) for item in result["row_issues"]}
        self.assertIn((3, "red"), messages)  # no price
        self.assertIn((4, "red"), messages)  # no MOQ
        self.assertIn((4, "yellow"), messages)  # no image
        self.assertIn((5, "yellow"), messages)  # duplicate SKU

    def test_rows_without_photos_stay_ready_and_one_photo_is_enough(self) -> None:
        book = Workbook()
        sheet = book.active
        sheet.append(["货号", "单价 USD", "起订量", "图片", "品名（中文）"])
        sheet.append(["A-01", "2.30", "300", "", "油漆刷"])
        sheet.append(["A-02", "1.00", "100", "only.jpg", "杯子"])
        buffer = io.BytesIO()
        book.save(buffer)

        result = preview(buffer.getvalue(), "simple")
        self.assertEqual(result["row_count"], 2)
        self.assertEqual(result["ready_count"], 2)
        self.assertEqual(result["blocked_count"], 0)
        self.assertEqual(result["image_stats"]["without_sheet_images"], 1)
        self.assertEqual(result["image_stats"]["partial_sheet_images"], 1)
        self.assertTrue(any("画一套" in item["message"] for item in result["row_issues"]))
        self.assertTrue(any("不满 6 张" in warning for warning in result["warnings"]))

        photos_only = preview(buffer.getvalue(), "simple", image_mode="photos_only")
        self.assertTrue(any("跳过" in item["message"] for item in photos_only["row_issues"]))

    def test_image_mode_decides_photos_generate_or_skip(self) -> None:
        self.assertEqual(decide_image_action(True, "mixed"), "use_photos")
        self.assertEqual(decide_image_action(False, "mixed"), "generate")
        self.assertEqual(decide_image_action(True, "generate_all"), "generate")
        self.assertEqual(decide_image_action(False, "photos_only"), "skip")
        self.assertEqual(decide_image_action(True, "photos_only"), "use_photos")
        self.assertEqual(decide_image_action(True, "complete_draw"), "complete")
        self.assertEqual(decide_image_action(True, "boost_skip"), "generate")
        self.assertEqual(decide_image_action(False, "complete_skip"), "skip")

    def test_brand_is_a_redline_default_not_an_ai_seed(self) -> None:
        row = ExcelRow(brand="Acme", price="1.80")
        self.assertEqual(row.extra_defaults(), {"brand": "Acme"})
        self.assertNotIn("brand", row.seed_values())

    def test_official_alibaba_template_only_asks_for_what_ai_cannot_know(self) -> None:
        payload = build_template(
            "alibaba",
            {"name": "油漆刷", "category_id": "21111112"},
            [{"id": "productTitle", "name": "Product name", "who": "AI 生成"}],
        )
        result = preview(payload, "alibaba")
        self.assertEqual(set(result["mapping"].values()), {"sku", "price", "moq", "images", "note"})
        self.assertNotIn("英文标题", result["headers"])

    def test_office_other_leaf_uses_stationery_family_not_generic(self) -> None:
        other = sheet_profile("Office & School Supplies / Other")
        self.assertEqual(other["family_id"], "stationery")
        labels = [item["label"] for item in other["spec_columns"]]
        self.assertIn("色数", labels)
        self.assertIn("硬度", labels)
        policy = fill_policy(
            [{"id": "leadHardness", "header": "Lead Hardness", "name": "Lead Hardness"}],
            other,
        )
        user_labels = [item["label"] for item in policy["user_fills"]]
        ai_labels = [item["label"] for item in policy["ai_fills"]]
        self.assertIn("色数", user_labels)
        self.assertNotIn("Lead Hardness", user_labels)
        self.assertTrue(any("硬度" in label or "Hardness" in label for label in ai_labels))
        self.assertIn("审过", policy["guarantee"])

    def test_category_sheet_adds_family_spec_columns_without_leaf_id(self) -> None:
        pencils = sheet_profile("Office & School Supplies / Colored Pencils")
        brushes = sheet_profile("Tools & Hardware / Paint Brushes")
        self.assertEqual(pencils["family_id"], "stationery")
        self.assertEqual(brushes["family_id"], "tools")
        pencil_labels = [item["label"] for item in pencils["spec_columns"]]
        brush_labels = [item["label"] for item in brushes["spec_columns"]]
        self.assertIn("色数", pencil_labels)
        self.assertNotIn("色数", brush_labels)
        self.assertIn("尺寸", brush_labels)
        payload = build_template("simple", category_name="Colored Pencils")
        book = load_workbook(io.BytesIO(payload))
        fill_sheet = next(name for name in book.sheetnames if name != "说明")
        headers = [cell.value for cell in next(book[fill_sheet].iter_rows(min_row=1, max_row=1))]
        self.assertIn("色数", headers)

    def test_leaf_category_puts_schema_attrs_on_fill_sheet(self) -> None:
        attrs = [
            {
                "id": "attr.icbuCatProp.p-2",
                "header": "Lead Hardness",
                "label": "铅芯硬度",
                "group": "icbuCatProp",
                "field_id": "p-2",
                "required": True,
                "options": [{"value": "HB", "label": "HB"}],
            }
        ]
        profile = sheet_profile("Colored Pencils", category_id="21110712", attr_columns=attrs)
        self.assertEqual(profile["family_id"], "leaf")
        preview_data = sheet_preview("simple", profile)
        labels = [item["label"] for item in preview_data["columns"]]
        self.assertIn("铅芯硬度", labels)
        payload = build_template(
            "simple",
            {"name": "Colored Pencils", "category_id": "21110712"},
            extra_columns=attrs,
            category_name="Colored Pencils",
            category_id="21110712",
        )
        book = load_workbook(io.BytesIO(payload))
        fill_sheet = next(name for name in book.sheetnames if name != "说明")
        headers = [cell.value for cell in next(book[fill_sheet].iter_rows(min_row=1, max_row=1))]
        self.assertIn("铅芯硬度", headers)
        self.assertNotIn("色数", headers)
        result = preview(payload, "simple", attrs)
        self.assertEqual(result["mapping"].get("铅芯硬度"), "attr.icbuCatProp.p-2")

    def test_spec_columns_parse_into_seller_facts(self) -> None:
        book = Workbook()
        sheet = book.active
        sheet.append(["货号", "单价 USD", "起订量", "图片", "色数", "硬度", "备注"])
        sheet.append(["CP-24", "2.40", "200", "a.jpg", "24", "HB", "水溶"])
        buffer = io.BytesIO()
        book.save(buffer)
        result = preview(buffer.getvalue(), "simple")
        rows = apply_preview(buffer.getvalue(), result["mapping"], "simple")
        self.assertEqual(rows[0].specs["color_count"], "24")
        self.assertEqual(rows[0].specs["hardness"], "HB")
        self.assertIn("色数 24", rows[0].fact_text())
        self.assertIn("HB", rows[0].fact_text())

    def test_dianxiaomi_template_stamps_the_listing_category(self) -> None:
        payload = build_template("dianxiaomi", {"name": "油漆刷", "category_id": "21111112"})
        result = preview(payload, "dianxiaomi")
        self.assertIn("英文标题", result["mapping"])
        self.assertNotIn("叶子类目 ID", result["headers"])


class FullSchemaTests(unittest.TestCase):
    SAMPLE = """<?xml version="1.0"?><itemSchema>
      <field id="productTitle" name="Product Title" type="input">
        <rules><rule name="requiredRule" value="true"/></rules>
      </field>
      <field id="icbuCatProp" type="complex"><fields>
        <field id="p-type" name="Type" type="singleCheck"><rules><rule name="requiredRule" value="true"/></rules>
          <options><option displayName="Oil Brush" value="1"/></options>
        </field>
        <field id="p-color" name="Color" type="singleCheck">
          <options><option displayName="Black" value="9"/></options>
        </field>
      </fields></field>
    </itemSchema>"""

    def test_schema_field_columns_include_optional(self) -> None:
        fields = parse_schema(self.SAMPLE)
        cols = schema_field_columns(fields)
        self.assertTrue(any(col["id"] == "schema.productTitle" and col["required"] for col in cols))
        self.assertTrue(any(col["id"] == "schema.icbuCatProp.p-type" and col["required"] for col in cols))
        self.assertTrue(any(col["id"] == "schema.icbuCatProp.p-color" and not col["required"] for col in cols))
        flat = flatten_schema_fields(fields)
        self.assertEqual(len(flat), len(cols))
        self.assertEqual(sum(1 for item in flat if item["required"]), 2)
        self.assertEqual(sum(1 for item in flat if not item["required"]), 1)

    def test_full_schema_profile_and_policy(self) -> None:
        fields = parse_schema(self.SAMPLE)
        extras = schema_field_columns(fields)
        profile = sheet_profile("Paint Brushes", category_id="21111112", attr_columns=extras)
        self.assertEqual(profile["family_id"], "full_schema")
        self.assertEqual(profile["required_count"], 2)
        self.assertEqual(profile["optional_count"], 1)
        policy = fill_policy(extras, profile)
        optional = [item for item in policy["user_fills"] if not item["required"]]
        self.assertTrue(any(item["id"] == "schema.icbuCatProp.p-color" for item in optional))

    def test_parse_schema_star_columns(self) -> None:
        extras = schema_field_columns(parse_schema(self.SAMPLE))
        rows = [
            ["货号", "单价 USD", "起订量", "英文标题", "类目属性 / 类型", "类目属性 / 颜色"],
            ["SKU-1", "1.8", "100", "Brush Set", "Oil Brush", "Black"],
        ]
        mapping = mapping_from_headers(rows[0], "full_schema", extras)
        parsed = parse_rows(rows, mapping, 0, extras)
        self.assertEqual(parsed[0].title, "Brush Set")
        self.assertEqual(parsed[0].attributes["icbuCatProp"]["p-type"], "1")
        self.assertEqual(parsed[0].attributes["icbuCatProp"]["p-color"], "9")


if __name__ == "__main__":
    unittest.main()

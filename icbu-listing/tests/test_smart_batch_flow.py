"""End-to-end smart batch flow up to review (parse grid, start concurrent image jobs)."""

import io
import json
import os
import sys
import tempfile
import unittest
import unittest.mock
import uuid
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "backend"))

os.environ["DATABASE_PATH"] = str(Path(tempfile.mkdtemp()) / f"smart-flow-{uuid.uuid4().hex}.db")
os.environ["UPLOAD_DIR"] = str(Path(tempfile.mkdtemp()) / "uploads")
os.environ["TOKEN_ENCRYPTION_KEY"] = "Frqkv73e3W19RfKb5WSgTueumPWcBSmpbM4wDyqu52o="
os.environ["REGISTRATION_CODES"] = "TEST-CODE"
os.environ["ALIBABA_APP_KEY"] = "test-key"
os.environ["ALIBABA_APP_SECRET"] = "test-secret"
os.environ.pop("OPENAI_API_KEY", None)

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from server.db import init_db  # noqa: E402
from server.main import app  # noqa: E402


def signup(tag: str = "user") -> TestClient:
    client = TestClient(app)
    email = f"{tag}-{uuid.uuid4().hex[:10]}@example.com"
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "supersecret", "code": "TEST-CODE"},
    )
    assert response.status_code == 200, response.text
    return client


def build_sample_xlsx(columns: list[dict]) -> bytes:
    book = Workbook()
    sheet = book.active
    labels = [col.get("label") or col["id"] for col in columns]
    for index, label in enumerate(labels, start=1):
        sheet.cell(1, index, label)
    values = {
        "sku": "FLOW-001",
        "price": "2.50",
        "moq": "300",
        "name": "测试画笔套装",
        "note": "12支装",
        "brand": "",
        "images": "",
    }
    for index, col in enumerate(columns, start=1):
        sheet.cell(2, index, "示例")
        sheet.cell(3, index, values.get(col["id"], "HB"))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class SmartBatchFlowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        init_db()

    def test_flow_to_review_with_concurrent_image_jobs(self) -> None:
        client = signup("smart-flow")
        shop_id = client.post("/api/v1/shops/bind-env", json={"name": "流程店"}).json()["id"]
        category_id = "21110712"

        with unittest.mock.patch(
            "server.services.smart_plan.catalog.get_schema_xml",
            return_value="<fields></fields>",
        ), unittest.mock.patch(
            "server.services.smart_plan.parse_schema",
            return_value=[],
        ):
            plan = client.get(
                "/api/v1/excel/smart-plan",
                params={"shop_id": shop_id, "category_id": category_id, "category_name": "彩铅"},
            )
        self.assertEqual(plan.status_code, 200, plan.text)
        columns = plan.json()["columns"]
        self.assertGreaterEqual(len(columns), 3)

        xlsx = build_sample_xlsx(columns)
        parse = client.post(
            "/api/v1/excel/doc-parse",
            data={
                "shop_id": shop_id,
                "category_id": category_id,
                "category_name": "彩铅",
                "image_mode": "complete_draw",
                "columns": json.dumps(columns, ensure_ascii=False),
            },
            files={"files": ("batch.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(parse.status_code, 200, parse.text)
        parsed = parse.json()
        self.assertGreaterEqual(parsed["row_count"], 1)
        self.assertGreaterEqual(len(parsed["rows"]), 1)
        self.assertEqual(parsed.get("planner"), "smart")

        started: list[str] = []

        def fake_start(user_id, row, *, category_id="", category_name=""):
            started.append(str(row.get("sku") or row.get("line")))
            return f"job-{len(started)}"

        with unittest.mock.patch(
            "server.routers.excel.grid_images.start_row_job",
            side_effect=fake_start,
        ), unittest.mock.patch(
            "server.routers.excel.grid_images.refresh_row_job",
            side_effect=lambda row, user_id: {**row, "image_job_status": "queued", "image_job_id": row.get("image_job_id", "job-x")},
        ):
            gen = client.post(
                "/api/v1/excel/grid-generate-images",
                data={
                    "shop_id": shop_id,
                    "category_id": category_id,
                    "category_name": "彩铅",
                    "rows": json.dumps(parsed["rows"], ensure_ascii=False),
                    "lines": "[]",
                },
            )
        self.assertEqual(gen.status_code, 200, gen.text)
        body = gen.json()
        self.assertGreaterEqual(len(started), parsed["row_count"])
        self.assertTrue(all(item.get("image_job_id") for item in body["rows"]))


if __name__ == "__main__":
    unittest.main()

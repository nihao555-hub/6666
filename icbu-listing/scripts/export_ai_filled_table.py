#!/usr/bin/env python3
"""Run live AI schema fill on sample fact rows and export a filled workbook.

Stops at AI fill — no publish. Output:
  samples/ai_filled_demo.xlsx  (填写 + AI填明细 + AI填宽表 + 质量摘要)
  samples/ai_filled_demo.json

Usage (from icbu-listing/):
  PYTHONPATH=backend:server python3 scripts/export_ai_filled_table.py
"""

from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "backend"))
sys.path.insert(0, str(ROOT / "server"))

from dotenv import load_dotenv  # noqa: E402

load_dotenv(ROOT / ".env")

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402

from ai import AiClient, ImageInput, Understanding  # noqa: E402
from schema import SchemaField, index_fields, parse_schema  # noqa: E402

from server.models import Shop  # noqa: E402
from server.services import catalog  # noqa: E402
from server.services.excel_import import SPEC_LABELS  # noqa: E402
from server.services.fact_bundle import FactBundle  # noqa: E402
from server.services.schema_fill import fill_category_draft  # noqa: E402
from server.services.shop_client import shop_api, shop_defaults  # noqa: E402

SAMPLES: list[dict[str, Any]] = [
    {
        "sku": "PEN-CP-12",
        "name": "12色木杆彩色铅笔",
        "category_id": "21110712",
        "price": "1.80",
        "moq": "500",
        "brand": "Giorgione",
        "note": "500@1.80; 1000@1.50; OEM logo welcome",
        "specs": {"color_count": "12", "material": "Wood", "hardness": "HB", "size": "17.5cm", "color": "colored"},
        "image_paths": sorted(str(p) for p in (ROOT / "data" / "generated-images" / "refs").glob("*.jpg"))[:4],
    },
    {
        "sku": "MKR-WB-24",
        "name": "双头美术马克笔24色",
        "category_id": "21110711",
        "price": "3.20",
        "moq": "200",
        "brand": "Giorgione",
        "note": "200@3.20; alcohol based; dual tip fine and chisel",
        "specs": {"color_count": "24", "material": "Plastic barrel", "tip": "Dual tip", "ink": "Alcohol"},
    },
    {
        "sku": "WC-PAN-18",
        "name": "18色固体水彩颜料",
        "category_id": "21111110",
        "price": "4.50",
        "moq": "300",
        "brand": "Giorgione",
        "note": "300@4.50; with brush; tin box",
        "specs": {"color_count": "18", "material": "Pigment", "packaging": "Tin box", "form": "Pan cake"},
    },
    {
        "sku": "SOFA-GD-01",
        "name": "户外花园PE藤编沙发三件套",
        "category_id": "100003017",
        "price": "89.00",
        "moq": "20",
        "brand": "",
        "note": "20 sets@89; UV resistant PE rattan; cushion included",
        "specs": {"material": "PE rattan", "color": "Brown", "pieces": "3", "frame": "Steel"},
    },
    {
        "sku": "FILTER-HEPA-01",
        "name": "空气净化器HEPA滤芯",
        "category_id": "100000050",
        "price": "6.50",
        "moq": "100",
        "brand": "",
        "note": "100@6.50; H13 HEPA; fits universal 352×478×30mm",
        "specs": {"material": "HEPA paper", "grade": "H13", "size": "352×478×30mm"},
    },
]

ATTR_GROUPS = ("icbuCatProp", "saleProp")


def _display(spec: SchemaField | None, raw: Any) -> str:
    if spec is None or raw in (None, "", [], {}):
        return ""
    if isinstance(raw, list):
        parts: list[str] = []
        for item in raw:
            token = str(item)
            option = next((o for o in spec.options if o.value == token), None)
            parts.append(option.display_name if option else token)
        return ", ".join(parts)
    token = str(raw)
    option = next((o for o in spec.options if o.value == token), None)
    if option is not None:
        return option.display_name
    return token


def _flatten_fields(
    specs: dict[str, SchemaField],
    values: dict[str, Any],
    evidence: dict[str, Any],
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for group_id in ATTR_GROUPS:
        group = specs.get(group_id)
        if group is None:
            continue
        group_values = values.get(group_id) or {}
        for child in group.children:
            raw = group_values.get(child.id)
            if raw in (None, "", [], {}):
                continue
            path = f"{group_id}.{child.id}"
            ev = evidence.get(path) or evidence.get(child.id) or {}
            rows.append(
                {
                    "group": group_id,
                    "field_id": child.id,
                    "field_name": child.name or child.id,
                    "value_raw": json.dumps(raw, ensure_ascii=False) if isinstance(raw, (list, dict)) else str(raw),
                    "value_label": _display(child, raw),
                    "required": "是" if child.required else "否",
                    "source": str(ev.get("source") or ""),
                    "quote": str(ev.get("quote") or "")[:120],
                }
            )
    for key in ("productTitle", "minOrderQuantity", "priceUnit", "scPrice"):
        spec = specs.get(key)
        raw = values.get(key)
        if raw in (None, "", [], {}):
            continue
        ev = evidence.get(key) or {}
        rows.append(
            {
                "group": "trade",
                "field_id": key,
                "field_name": spec.name if spec else key,
                "value_raw": str(raw),
                "value_label": _display(spec, raw) if spec else str(raw),
                "required": "是" if spec and spec.required else "否",
                "source": str(ev.get("source") or "excel"),
                "quote": str(ev.get("quote") or "")[:120],
            }
        )
    ladder = values.get("ladderPrice") or {}
    if isinstance(ladder, dict):
        for slot, payload in ladder.items():
            if not isinstance(payload, dict):
                continue
            ev = evidence.get(f"ladderPrice.{slot}") or {}
            rows.append(
                {
                    "group": "trade",
                    "field_id": slot,
                    "field_name": "阶梯价",
                    "value_raw": json.dumps(payload, ensure_ascii=False),
                    "value_label": f"{payload.get('quantity')} @ {payload.get('price')}",
                    "required": "是",
                    "source": str(ev.get("source") or "excel"),
                    "quote": str(ev.get("quote") or "")[:120],
                }
            )
    return rows


def _style_header(sheet, row: int = 1, color: str = "1D4ED8") -> None:
    fill = PatternFill("solid", fgColor=color)
    font = Font(color="FFFFFF", bold=True)
    for cell in sheet[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(sheet, max_width: int = 48) -> None:
    for col in range(1, sheet.max_column + 1):
        letter = get_column_letter(col)
        width = 10
        for row in range(1, min(sheet.max_row, 200) + 1):
            value = sheet.cell(row, col).value
            if value:
                width = max(width, min(max_width, len(str(value)) + 2))
        sheet.column_dimensions[letter].width = width


def build_workbook(results: list[dict[str, Any]]) -> bytes:
    book = Workbook()

    # --- 填写表 ---
    fill = book.active
    fill.title = "填写"
    fill_headers = [
        "货号",
        "品名",
        "叶子类目ID",
        "单价USD",
        "起订量",
        "品牌",
        "备注",
        "规格列(JSON)",
    ]
    fill.append(fill_headers)
    for item in results:
        sample = item["sample"]
        fill.append(
            [
                sample["sku"],
                sample["name"],
                sample["category_id"],
                sample["price"],
                sample["moq"],
                sample.get("brand", ""),
                sample.get("note", ""),
                json.dumps(sample.get("specs") or {}, ensure_ascii=False),
            ]
        )
    _style_header(fill, color="171717")

    # --- AI填明细 ---
    detail = book.create_sheet("AI填-明细")
    detail_headers = [
        "货号",
        "类目",
        "字段组",
        "字段ID",
        "字段名",
        "填写值(平台值)",
        "显示标签",
        "必填",
        "来源",
        "依据摘录",
        "状态",
    ]
    detail.append(detail_headers)
    for item in results:
        sample = item["sample"]
        blocked = {issue.get("path") for issue in item.get("issues") or [] if issue.get("level") == "red"}
        for field in item["fields"]:
            path = f"{field['group']}.{field['field_id']}" if field["group"] in ATTR_GROUPS else field["field_id"]
            status = "已填"
            if path in blocked or field["field_id"] in blocked:
                status = "缺依据/待人工"
            detail.append(
                [
                    sample["sku"],
                    item["category_label"],
                    field["group"],
                    field["field_id"],
                    field["field_name"],
                    field["value_raw"],
                    field["value_label"],
                    field["required"],
                    field["source"],
                    field["quote"],
                    status,
                ]
            )
        for issue in item.get("issues") or []:
            if issue.get("level") != "red":
                continue
            detail.append(
                [
                    sample["sku"],
                    item["category_label"],
                    "issue",
                    issue.get("field_id", ""),
                    issue.get("field_name", ""),
                    "",
                    "",
                    "是" if "必填" in str(issue.get("message", "")) else "否",
                    "blocked",
                    str(issue.get("message", ""))[:120],
                    "未填-待人工",
                ]
            )
    _style_header(detail)

    # --- AI填宽表 ---
    wide = book.create_sheet("AI填-宽表")
    wide_keys: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for item in results:
        for field in item["fields"]:
            key = (field["group"], field["field_id"])
            if key not in seen:
                seen.add(key)
                wide_keys.append(key)
    wide_header = ["货号", "类目", "标题", "MOQ", "阶梯价"] + [f"{g}.{fid}" for g, fid in wide_keys]
    wide.append(wide_header)
    for item in results:
        sample = item["sample"]
        values = item["values"]
        ladder = values.get("ladderPrice") or {}
        ladder_text = "; ".join(
            f"{payload.get('quantity')}@{payload.get('price')}"
            for payload in ladder.values()
            if isinstance(payload, dict)
        )
        lookup = {(f["group"], f["field_id"]): f["value_label"] for f in item["fields"]}
        row = [
            sample["sku"],
            item["category_label"],
            values.get("productTitle", ""),
            values.get("minOrderQuantity", sample["moq"]),
            ladder_text,
        ]
        row.extend(lookup.get(key, "") for key in wide_keys)
        wide.append(row)
    _style_header(wide, color="047857")

    # --- 质量摘要 ---
    summary = book.create_sheet("质量摘要")
    summary.append(
        [
            "货号",
            "类目",
            "模型",
            "AI调用次数",
            "已填属性数",
            "AI填数",
            "本地/店铺填数",
            "红色问题数",
            "必填选项总数",
            "已填必填选项",
            "仍缺必填",
            "说明",
        ]
    )
    for item in results:
        stats = item["stats"]
        summary.append(
            [
                item["sample"]["sku"],
                item["category_label"],
                item.get("model", ""),
                stats.get("ai_calls", 0),
                len(item["fields"]),
                stats.get("filled_ai", 0),
                stats.get("filled_local", 0),
                stats.get("red_issues", 0),
                stats.get("required_option_fields", 0),
                stats.get("filled_required_options", 0),
                stats.get("missing_required_options", 0),
                item.get("note", ""),
            ]
        )
    summary.append([])
    summary.append(["生成时间(UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    summary.append(["流程", "Excel事实+商品图 → vision识图 → AI填官方属性 → 导出（未发布）"])
    _style_header(summary, color="7C3AED")

    for sheet in book.worksheets:
        _autosize(sheet)

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _load_images(sample: dict[str, Any]) -> list[ImageInput]:
    out: list[ImageInput] = []
    for raw in sample.get("image_paths") or []:
        path = Path(raw)
        if not path.is_file():
            continue
        out.append(ImageInput(filename=path.name, content=path.read_bytes()))
    return out[:6]


def run_one(
    db,
    api,
    ai: AiClient,
    sample: dict[str, Any],
    defaults: dict[str, Any],
) -> dict[str, Any]:
    category_id = str(sample["category_id"])
    node = catalog.get_node(db, api, category_id, fetch=False)
    label = catalog.label(node) if node else category_id
    xml = catalog.get_schema_xml(db, api, category_id, str(defaults.get("language") or "en_US"), fetch=False)
    fields = parse_schema(xml)
    specs = index_fields(fields)

    bundle = FactBundle(
        sku=sample["sku"],
        name=sample["name"],
        brand=sample.get("brand") or "",
        note=sample.get("note") or "",
        specs={k: str(v) for k, v in (sample.get("specs") or {}).items()},
        price=sample["price"],
        moq=sample["moq"],
        origin="China",
    )
    images = _load_images(sample)
    understanding = Understanding(
        product_name=sample["name"],
        material=bundle.specs.get("material", ""),
        colors=[part.strip() for part in str(bundle.specs.get("color", "")).split(",") if part.strip()],
        usage=bundle.specs.get("usage", ""),
        features=["OEM welcome", "Bulk order"],
        specs=dict(bundle.specs),
    )
    if images and ai is not None:
        try:
            understanding = bundle.enrich(ai.understand(images, bundle.text_blob() or sample["name"]))
        except Exception:
            understanding = bundle.enrich(understanding)
    else:
        understanding = bundle.enrich(understanding)
    report = fill_category_draft(
        xml,
        understanding=understanding,
        bundle=bundle,
        defaults=defaults,
        category_values={},
        price=sample["price"],
        moq=sample["moq"],
        ai=ai,
        images_applied=bool(images),
        category_id=category_id,
        images=images,
        title=f"Wholesale colored pencil set",
        keywords=["wholesale", "factory", "bulk"],
        highlights=bundle.text_blob(),
    )
    evidence = {path: ev.as_dict() for path, ev in report.evidence.items()}
    field_rows = _flatten_fields(specs, report.values, evidence)

    required_option_fields = 0
    filled_required_options = 0
    filled_ai = 0
    filled_local = 0
    for group_id in ATTR_GROUPS:
        group = specs.get(group_id)
        if group is None:
            continue
        group_values = report.values.get(group_id) or {}
        for child in group.children:
            if not child.required or not child.options:
                continue
            required_option_fields += 1
            if child.id not in group_values:
                continue
            filled_required_options += 1
            path = f"{group_id}.{child.id}"
            source = (evidence.get(path) or evidence.get(child.id) or {}).get("source")
            if source == "ai":
                filled_ai += 1
            else:
                filled_local += 1

    red_issues = [issue for issue in report.issues if issue.get("level") == "red"]
    return {
        "sample": sample,
        "category_label": label,
        "values": report.values,
        "evidence": evidence,
        "issues": report.issues,
        "fields": field_rows,
        "stats": {
            **report.stats.as_dict(),
            "filled_ai": filled_ai,
            "filled_local": filled_local,
            "required_option_fields": required_option_fields,
            "filled_required_options": filled_required_options,
            "missing_required_options": required_option_fields - filled_required_options,
            "red_issues": len(red_issues),
        },
        "model": ai.text_model,
        "note": "无依据的必填项留空标红；有图时会先 vision 识图再 multimodal 填属性",
    }


def main() -> int:
    ai = AiClient.from_env()
    print(f"model={ai.text_model} base={ai.base_url}")

    engine = create_engine(f"sqlite:///{ROOT / 'data' / 'auto-shoper.db'}")
    Session = sessionmaker(bind=engine)
    db = Session()
    shop = db.query(Shop).filter(Shop.access_token != "").first()
    if shop is None:
        print("No shop token in database.", file=sys.stderr)
        return 1
    api = shop_api(shop)
    defaults = dict(shop_defaults(shop))
    defaults.setdefault("origin", "China")
    defaults.setdefault("priceUnit", "Piece/Pieces")
    defaults.setdefault("saleType", "Unit")
    defaults.setdefault("ladderPeriod", "15")
    defaults.setdefault("paymentMethod", "T/T")
    defaults.setdefault("port", "Ningbo")
    defaults.setdefault("market", "询盘")
    defaults.setdefault("logisticsProperty", "general_cargo_0")
    defaults.setdefault("marketSample", "Unavailable")
    defaults.setdefault("language", "en_US")

    results: list[dict[str, Any]] = []
    for index, sample in enumerate(SAMPLES, start=1):
        print(f"[{index}/{len(SAMPLES)}] {sample['sku']} cat={sample['category_id']} …", flush=True)
        try:
            results.append(run_one(db, api, ai, sample, defaults))
            stats = results[-1]["stats"]
            print(
                f"  filled={len(results[-1]['fields'])} ai={stats['filled_ai']} "
                f"local={stats['filled_local']} red={stats['red_issues']}",
                flush=True,
            )
        except Exception as exc:  # noqa: BLE001
            print(f"  FAILED: {exc}", flush=True)
            results.append(
                {
                    "sample": sample,
                    "category_label": sample["category_id"],
                    "values": {},
                    "evidence": {},
                    "issues": [{"level": "red", "message": str(exc)[:200]}],
                    "fields": [],
                    "stats": {"red_issues": 1},
                    "model": ai.text_model,
                    "note": "失败",
                }
            )

    db.close()

    out_dir = ROOT / "samples"
    out_dir.mkdir(exist_ok=True)
    xlsx_path = out_dir / "ai_filled_demo.xlsx"
    json_path = out_dir / "ai_filled_demo.json"
    xlsx_path.write_bytes(build_workbook(results))
    json_path.write_text(
        json.dumps(
            {
                "generated_at": datetime.now(timezone.utc).isoformat(),
                "model": ai.text_model,
                "rows": results,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"wrote {xlsx_path}")
    print(f"wrote {json_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
